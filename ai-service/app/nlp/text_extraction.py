"""Document text extraction for the resume pipeline (§8).

Supports PDF (PyMuPDF), DOCX (python-docx) and plain text. Both libraries are
imported lazily so the service starts even if one is missing, and the failure is
reported as a clear, user-facing message rather than an import crash.
"""
from __future__ import annotations

import io
import re

from app.core.errors import ResumeParseError, UnsupportedFileType
from app.core.logging import get_logger

log = get_logger(__name__)

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DOC_MIME = "application/msword"
TXT_MIME = "text/plain"

SUPPORTED_MIMES = {PDF_MIME, DOCX_MIME, DOC_MIME, TXT_MIME}


def _clean(text: str) -> str:
    """Normalises whitespace and repairs the artefacts PDF extraction leaves.

    PDFs commonly yield ligatures, non-breaking spaces, hyphenated line breaks
    and bullet glyphs that would otherwise corrupt keyword matching.
    """
    if not text:
        return ""

    replacements = {
        " ": " ", "’": "'", "‘": "'", "“": '"', "”": '"',
        "–": "-", "—": "-", "ﬁ": "fi", "ﬂ": "fl", "•": "\n- ",
        "●": "\n- ", "▪": "\n- ", "·": "\n- ", "": "\n- ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)

    # Re-join words split across a line break by hyphenation.
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    # Drop control characters but keep newlines and tabs.
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    # Collapse runs of spaces, and runs of blank lines down to one.
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    return text.strip()


def _extract_pdf(data: bytes) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise ResumeParseError(
            "PDF support is unavailable because PyMuPDF is not installed. "
            "Run: python -m pip install PyMuPDF"
        ) from exc

    try:
        with fitz.open(stream=data, filetype="pdf") as doc:
            if doc.is_encrypted and not doc.authenticate(""):
                raise ResumeParseError(
                    "This PDF is password protected. Please upload an unprotected copy of your resume."
                )
            pages = [page.get_text("text") for page in doc]
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError(
            "The PDF could not be read. It may be corrupted - try re-exporting it from your editor."
        ) from exc

    return _clean("\n".join(pages))


def _extract_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise ResumeParseError(
            "DOCX support is unavailable because python-docx is not installed. "
            "Run: python -m pip install python-docx"
        ) from exc

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeParseError(
            "The Word document could not be read. If it is an older .doc file, "
            "please save it as .docx or PDF and upload again."
        ) from exc

    parts = [p.text for p in document.paragraphs]
    # Many resumes lay content out in tables; skipping those loses whole sections.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return _clean("\n".join(parts))


def _extract_txt(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return _clean(data.decode(encoding))
        except UnicodeDecodeError:
            continue
    raise ResumeParseError("The text file could not be decoded. Please save it as UTF-8 and try again.")


def extract_text(data: bytes, mime_type: str, filename: str = "") -> str:
    """Extracts plain text from an uploaded resume.

    Raises `ResumeParseError` when the file yields too little text to analyse -
    the usual cause is a scanned image PDF, which this build does not OCR.
    """
    if not data:
        raise ResumeParseError("The uploaded file is empty.")

    mime = (mime_type or "").lower()
    lower_name = filename.lower()

    if mime == PDF_MIME or lower_name.endswith(".pdf"):
        text = _extract_pdf(data)
    elif mime == DOCX_MIME or lower_name.endswith(".docx"):
        text = _extract_docx(data)
    elif mime == TXT_MIME or lower_name.endswith(".txt"):
        text = _extract_txt(data)
    elif mime == DOC_MIME or lower_name.endswith(".doc"):
        raise UnsupportedFileType(
            "Legacy .doc files are not supported. Please save your resume as .docx or PDF and upload it again."
        )
    else:
        raise UnsupportedFileType(
            f"'{mime_type or filename}' is not a supported resume format. Upload a PDF, DOCX or TXT file."
        )

    if len(text) < 120:
        raise ResumeParseError(
            "Almost no readable text could be extracted from this file. "
            "If your resume is a scanned image, please upload a text-based PDF or a DOCX file instead."
        )

    log.info("Resume text extracted", extra={"chars": len(text), "mime": mime})
    return text
